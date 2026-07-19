
import pandas as pd, re, io, smtplib, gzip, base64, json, sys
import gspread
from google.oauth2.service_account import Credentials
import toml
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# order_id передаётся аргументом: python3 resend.py <order_id> [email_override]
ORDER_ID = sys.argv[1]
EMAIL_OVERRIDE = sys.argv[2] if len(sys.argv) > 2 else None

EMAIL_FROM = "result@vuzline.ru"
EMAIL_PASSWORD = "kpamwwoldfbqaooh"

secrets = toml.load(".streamlit/secrets.toml")
SHEETS_ID = secrets["SHEETS_ID"]
creds = Credentials.from_service_account_info(
    dict(secrets["gcp_service_account"]),
    scopes=["https://spreadsheets.google.com/feeds", "https://www.googleapis.com/auth/drive"]
)
client = gspread.authorize(creds)
sheet = client.open_by_key(SHEETS_ID).sheet1

cell = sheet.find(ORDER_ID, in_column=1)
if not cell:
    print(f"Order {ORDER_ID} не найден в Sheets")
    sys.exit(1)

row_values = sheet.row_values(cell.row)
user_email = EMAIL_OVERRIDE or row_values[2]
user_email = "".join(c for c in str(user_email).strip() if ord(c) < 128)
compressed_b64 = "".join(row_values[6:])
raw_json = gzip.decompress(base64.b64decode(compressed_b64)).decode("utf-8")
data = json.loads(raw_json)
result_df = pd.DataFrame.from_dict(data["result"])
search_params = data["search_params"]
print(f"Данные найдены: {len(result_df)} строк, email={user_email}")

for col in result_df.columns:
    result_df[col] = result_df[col].apply(
        lambda x: "" if x is None or (isinstance(x, float) and pd.isna(x)) or str(x) in ("None", "nan") else x
    )

CC={"Уверенно":"D6EED2","Реалистично":"D0E8F5","Вероятно":"FFF0D6","Рискованно":"FAE0E0","Маловероятно":"EBEBEB","Квоты и БВИ":"D7EAF3","Общего конкурса не было":"D8D8D8","Нет данных":"F5F5F5","Нет оценки — не указан балл за ДВИ":"F5F5F5"}
CT={"Уверенно":"1E6B14","Реалистично":"0D5A8A","Вероятно":"8A5A00","Рискованно":"8A1A1A","Маловероятно":"555555","Квоты и БВИ":"1A5C7A","Общего конкурса не было":"4A4A4A","Нет данных":"888888","Нет оценки — не указан балл за ДВИ":"888888"}
HB="379FFC"; HF="FFFFFF"; RO="FFFFFF"; RE="F4F7FB"; AC="379FFC"
thin_g=Side(style="thin",color="DDDDDD"); no_s=Side(style=None); dashed=Side(style="dashDot",color="8BB8E8")

def gborder(ci,bottom=True):
    b=thin_g if bottom else no_s
    l=dashed if ci==6 else no_s
    r=dashed if ci in [1,6,8] else no_s
    return Border(bottom=b,left=l,right=r)

wb=Workbook()
ws3=wb.active; ws3.title="Важная информация"; ws3.column_dimensions["A"].width=90
info=[("📋 Результаты во вкладке Результаты","title"),("Ниже — как читать таблицу и расставлять приоритеты.","subtitle"),("","gap"),("О СЕРВИСЕ","header"),("Данные актуальны на 2026 год. Носят рекомендательный характер.","text"),("Только для поступающих на общих основаниях на основном этапе.","text"),("","gap"),("РАСШИФРОВКА ШАНСОВ","header"),("Уверенно — балл выше среднего на 10+. Рек. приоритет: 3–5+","D6EED2"),("Реалистично — балл выше среднего. Рек. приоритет: 2–3","D0E8F5"),("Вероятно — балл ниже среднего, но выше проходного на 5+. Рек. приоритет: 1–2","FFF0D6"),("Рискованно — балл близко к проходному. Рек. приоритет: 1*","FAE0E0"),("Маловероятно — балл ниже проходного на 15+.","EBEBEB"),("Квоты и БВИ — места для общего конкурса могут не остаться. Известно 3 августа.","D7EAF3"),("Нет данных — новая специальность.","F5F5F5"),("","gap"),("СРОКИ","header"),("До 25 июля — подача документов.","text"),("27 июля — конкурсные списки.","text"),("5 августа 12:00 мск — последний срок согласия на зачисление.","text"),("","gap"),("Консультация: https://vuzline.ru/consultation | result@vuzline.ru","contact")]
for rn,(text,kind) in enumerate(info,1):
    cell_=ws3.cell(row=rn,column=1,value=text)
    cell_.alignment=Alignment(wrap_text=True,vertical="center")
    if kind=="title": cell_.font=Font(name="Montserrat",size=16,bold=True,color=HB); ws3.row_dimensions[rn].height=42
    elif kind=="subtitle": cell_.font=Font(name="Montserrat",size=10,color="555555"); ws3.row_dimensions[rn].height=22
    elif kind=="gap": ws3.row_dimensions[rn].height=10
    elif kind=="header": cell_.font=Font(name="Montserrat",size=10,bold=True,color=HB); cell_.fill=PatternFill("solid",fgColor="EBF4FF"); ws3.row_dimensions[rn].height=24
    elif kind=="contact": cell_.font=Font(name="Montserrat",size=9,bold=True,color=AC); ws3.row_dimensions[rn].height=36
    elif len(kind)==6: cell_.font=Font(name="Montserrat",size=9); cell_.fill=PatternFill("solid",fgColor=kind); ws3.row_dimensions[rn].height=20
    else: cell_.font=Font(name="Montserrat",size=9); ws3.row_dimensions[rn].height=18

ws=wb.create_sheet("Результаты")
df_out=result_df.copy()
cl={"podstrahovka":"Уверенно","realistic":"Реалистично","probable":"Вероятно","risky":"Рискованно","unlikely":"Маловероятно","quota_bvi":"Квоты и БВИ","no_competition":"Общего конкурса не было","new":"Нет данных","no_dvi_score":"Нет оценки — не указан балл за ДВИ"}
if "Шансы" in df_out.columns: df_out["Шансы"]=df_out["Шансы"].map(lambda x:cl.get(x,x))
cols=list(df_out.columns)
for ci,cn in enumerate(cols,1):
    c=ws.cell(row=1,column=ci,value=cn)
    c.font=Font(bold=True,color=HF,name="Montserrat",size=9)
    c.fill=PatternFill("solid",fgColor=HB)
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    c.border=gborder(ci,bottom=False)
ws.row_dimensions[1].height=32
chance_col_idx = cols.index("Шансы")+1 if "Шансы" in cols else -1
prio_col_idx = cols.index("Рек. приоритет")+1 if "Рек. приоритет" in cols else -1
for ri,row_ in enumerate(df_out.itertuples(index=False),start=2):
    row_=pd.Series(row_._asdict())
    bg=RO if ri%2==0 else RE
    for ci,(cn,value) in enumerate(row_.items(),1):
        if value is None or (isinstance(value,float) and pd.isna(value)): cv=""
        elif isinstance(value,(int,float)) and not isinstance(value,bool): cv=value
        else:
            s=str(value); cv="" if s in ("None","nan") else s
        c=ws.cell(row=ri,column=ci,value=cv)
        c.fill=PatternFill("solid",fgColor=bg)
        c.border=gborder(ci)
        if ci in [1,2,3,4,5]: c.font=Font(name="Montserrat",size=9,bold=(ci==1)); c.alignment=Alignment(horizontal="left",vertical="center")
        elif ci==chance_col_idx:
            cc=CC.get(str(cv),bg); tt=CT.get(str(cv),"1A1A1A")
            c.fill=PatternFill("solid",fgColor=cc); c.font=Font(name="Montserrat",size=9,color=tt,bold=True); c.alignment=Alignment(horizontal="center",vertical="center")
        elif ci==prio_col_idx: c.font=Font(name="Montserrat",size=9,color=AC,bold=True); c.alignment=Alignment(horizontal="center",vertical="center")
        else: c.font=Font(name="Montserrat",size=9); c.alignment=Alignment(horizontal="center",vertical="center")
ws.freeze_panes="A2"
cw={1:14,2:18,3:22,4:32,5:30,6:7,7:13,8:13,9:13,10:10,11:13,12:18,13:11}
for ci,w in cw.items(): ws.column_dimensions[get_column_letter(ci)].width=w

ws2=wb.create_sheet("Запрос")
for ci,(k,v) in enumerate(search_params.items(),1):
    ws2.cell(row=1,column=ci,value=k).font=Font(bold=True,color=HF,name="Montserrat",size=9)
    ws2.cell(row=1,column=ci).fill=PatternFill("solid",fgColor=HB)
    ws2.cell(row=2,column=ci,value=str(v)).font=Font(name="Montserrat",size=9)
    ws2.column_dimensions[get_column_letter(ci)].width=25

buf=io.BytesIO(); wb.save(buf); buf.seek(0)
msg=MIMEMultipart()
msg["From"]=EMAIL_FROM; msg["To"]=user_email; msg["Subject"]="Ваша таблица подбора вузов — Vuzline"
body="Здравствуйте!\n\nВаша таблица готова — она прикреплена к этому письму.\n\nИзвините за задержку — произошёл технический сбой.\n\nУдачи с поступлением!\nКоманда Vuzline\nresult@vuzline.ru"
msg.attach(MIMEText(body,"plain","utf-8"))
att=MIMEBase("application","octet-stream"); att.set_payload(buf.read()); encoders.encode_base64(att)
att.add_header("Content-Disposition","attachment; filename=\"vuzline_results.xlsx\"")
msg.attach(att)
print(f"Отправляем на {user_email}...")
with smtplib.SMTP_SSL("smtp.yandex.ru",465) as s:
    s.login(EMAIL_FROM,EMAIL_PASSWORD); s.sendmail(EMAIL_FROM,user_email,msg.as_string())
sheet.update_cell(cell.row, 6, "sent")
print("Готово! Письмо отправлено, статус sent записан.")
